from __future__ import annotations

from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
import json
import os
import re
import uuid

from mcp.server import MCPServer
from mcp.server.transport_security import TransportSecuritySettings
from starlette.applications import Starlette
from starlette.responses import FileResponse, JSONResponse, PlainTextResponse
from starlette.routing import Mount, Route

from motor_docx import fill_docx, collect_placeholders
from docx import Document

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE = BASE_DIR / 'templates' / 'MASTER_SRC_TEMPLATE.docx'
FIELD_MAP = BASE_DIR / 'config' / 'field_map.json'
GENERATED_DIR = Path(os.getenv('GENERATED_DIR', str(BASE_DIR / 'generated')))
GENERATED_DIR.mkdir(parents=True, exist_ok=True)
PUBLIC_BASE_URL = os.getenv('PUBLIC_BASE_URL', '').rstrip('/')

mcp = MCPServer(
    'SRC Sesiones Quincenales',
    version='0.1.1',
    instructions=(
        'Genera documentos DOCX de sesiones quincenales usando exclusivamente la plantilla institucional SRC. '
        'No rediseñes la plantilla. Antes de generar, construye contenido pedagógico detallado para Inicio, Desarrollo '
        'y Cierre. Usa negrita solo en etiquetas/subtítulos; los párrafos explicativos van en texto normal y separados.'
    ),
)


def _field_keys() -> list[str]:
    """
    Read the real required fields directly from MASTER_SRC_TEMPLATE.docx.

    This deliberately does NOT infer fields from field_map.json because that
    file is descriptive/grouped metadata, not the authoritative source of
    placeholder names.
    """
    if not TEMPLATE.exists():
        return []
    return collect_placeholders(Document(TEMPLATE))


def _field_groups() -> dict:
    """Return descriptive grouping metadata for ChatGPT, when available."""
    try:
        raw = json.loads(FIELD_MAP.read_text(encoding='utf-8'))
        return raw if isinstance(raw, dict) else {}
    except Exception:
        return {}


def _safe_docx_name(name: str) -> str:
    name = (name or 'SESIONES_QUINCENALES.docx').strip()
    if not name.lower().endswith('.docx'):
        name += '.docx'
    stem = Path(name).stem
    stem = re.sub(r'[^A-Za-z0-9ÁÉÍÓÚÜÑáéíóúüñ._ -]+', '_', stem).strip(' ._')
    return (stem or 'SESIONES_QUINCENALES') + '.docx'


def _find_generated(file_id: str) -> Path | None:
    for path in GENERATED_DIR.glob(f'{file_id}__*.docx'):
        if path.is_file():
            return path
    return None


@mcp.tool()
def get_template_info() -> dict:
    """Consulta la plantilla SRC y los campos que el complemento puede rellenar. No modifica archivos."""
    fields = _field_keys()
    return {
        'template': TEMPLATE.name,
        'version': '0.1.1',
        'field_count': len(fields),
        'fields': fields,
        'field_groups': _field_groups(),
        'important_instruction': (
            'Before calling generate_quincena_docx, prepare ALL fields returned '
            'in fields. The generator rejects incomplete payloads instead of '
            'creating a blank Word.'
        ),
        'format_rules': [
            'Conservar diseño, logo, pie, tablas, colores y tipografía de la plantilla.',
            'Negrita solo en etiquetas y subtítulos.',
            'Inicio, Desarrollo y Cierre deben estar desarrollados al detalle.',
            'Cada subapartado pedagógico debe ser un párrafo Word independiente con separación visual.',
        ],
        'current_limitation': 'La plantilla prototipo actual dispone de 6 bloques de sesión; la clonación dinámica será la siguiente versión.',
    }


@mcp.tool()
def generate_quincena_docx(fields: dict[str, str], output_name: str = 'SESIONES_QUINCENALES_SRC.docx') -> dict:
    """
    Genera un Word institucional de sesiones quincenales SRC a partir de campos ya redactados.
    Usa esta herramienta después de que ChatGPT haya preparado el contenido pedagógico completo.
    """
    required = _field_keys()
    missing = [key for key in required if key not in fields]
    if missing:
        return {
            'ok': False,
            'error': 'Faltan campos requeridos por la plantilla. No se generó ningún Word.',
            'missing_fields': missing,
            'required_field_count': len(required),
            'received_field_count': len(fields),
        }

    # Reject an obviously empty/placeholder payload even when every key exists.
    critical = [
        'UNIDAD', 'TITULO_UNIDAD', 'NIVEL', 'AREA', 'GRADO', 'DOCENTE',
        'SITUACION_CONTEXTO', 'COMPETENCIAS',
        'SESION_1_INICIO', 'SESION_1_DESARROLLO', 'SESION_1_CIERRE'
    ]
    empty_critical = [
        key for key in critical
        if not str(fields.get(key, '')).strip()
    ]
    if empty_critical:
        return {
            'ok': False,
            'error': 'Hay campos críticos vacíos. No se generó ningún Word.',
            'empty_critical_fields': empty_critical,
        }

    file_id = uuid.uuid4().hex
    safe_name = _safe_docx_name(output_name)
    path = GENERATED_DIR / f'{file_id}__{safe_name}'
    try:
        fill_docx(TEMPLATE, fields, path)
    except Exception as exc:
        return {'ok': False, 'error': str(exc)}

    relative = f'/download/{file_id}'
    return {
        'ok': True,
        'file_id': file_id,
        'filename': safe_name,
        'download_url': f'{PUBLIC_BASE_URL}{relative}' if PUBLIC_BASE_URL else relative,
        'message': 'Documento DOCX generado correctamente con la plantilla institucional SRC.',
    }


@mcp.tool()
def get_generated_document(file_id: str) -> dict:
    """Consulta si un documento generado todavía está disponible para descargar. No modifica archivos."""
    path = _find_generated(file_id)
    if path is None:
        return {'found': False, 'file_id': file_id}
    relative = f'/download/{file_id}'
    original_name = path.name.split('__', 1)[1] if '__' in path.name else path.name
    return {
        'found': True,
        'file_id': file_id,
        'filename': original_name,
        'size_bytes': path.stat().st_size,
        'download_url': f'{PUBLIC_BASE_URL}{relative}' if PUBLIC_BASE_URL else relative,
    }


async def health(request):
    return JSONResponse({
        'ok': True,
        'service': 'SRC Sesiones Quincenales',
        'version': '0.1.1',
        'mcp_endpoint': '/mcp',
    })


async def home(request):
    return PlainTextResponse('SRC Sesiones Quincenales MCP v0.1.1 - endpoint: /mcp')


async def download(request):
    file_id = request.path_params['file_id']
    path = _find_generated(file_id)
    if path is None:
        return PlainTextResponse('Documento no encontrado o expirado.', status_code=404)
    original_name = path.name.split('__', 1)[1] if '__' in path.name else path.name
    return FileResponse(path, filename=original_name, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


# Render/other reverse proxies already control the public Host header. For local dev,
# leave the protection on; for deployment set MCP_DISABLE_DNS_REBINDING=true.
if os.getenv('MCP_DISABLE_DNS_REBINDING', '').lower() in {'1', 'true', 'yes'}:
    security = TransportSecuritySettings(enable_dns_rebinding_protection=False)
else:
    security = None

mcp_app = mcp.streamable_http_app(json_response=True, transport_security=security)


@asynccontextmanager
async def lifespan(app: Starlette) -> AsyncIterator[None]:
    async with mcp.session_manager.run():
        yield


app = Starlette(
    routes=[
        Route('/', home, methods=['GET']),
        Route('/health', health, methods=['GET']),
        Route('/download/{file_id}', download, methods=['GET']),
        Mount('/', app=mcp_app),
    ],
    lifespan=lifespan,
)
