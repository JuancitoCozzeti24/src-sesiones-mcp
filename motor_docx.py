from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import json
import re
from typing import Mapping

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt

BOLD_LABELS = {
    'Fecha sugerida','Fecha','Grado y sección','Grado','Tema','Competencia','Capacidades',
    'Desempeño precisado','Criterio de evaluación','Criterios de evaluación','Propósito de aprendizaje','Propósito',
    'Motivación','Saberes previos','Problematización','Problematización / conflicto cognitivo','Conflicto cognitivo',
    'Declaración del propósito y acuerdos','Propósito y acuerdos','Construcción del conocimiento','Exploración',
    'Modelado docente','Modelado','Práctica guiada','Práctica autónoma','DUA','DUA y acompañamiento',
    'Atención a la diversidad / DUA','Atención a condiciones de acceso','Comparación directa-inversa',
    'Comparación de procedimientos','Análisis de errores','Evaluación formativa','Práctica autónoma y evaluación formativa',
    'Evaluación individual','Monitoreo docente','Revisión final del estudiante','Registro de evidencias',
    'Retroalimentación','Retroalimentación y síntesis','Síntesis','Síntesis del procedimiento','Metacognición','Transferencia',
    'Ticket de salida','Cierre formativo','Proyección','Activación inicial','Orientaciones','Clima de evaluación',
    'Gamificación','Gamificación / Dinámica','Socialización','Formalización','Acompañamiento docente'
}

# These labels start pedagogical blocks. They get visible paragraph spacing.
SECTION_LABELS = BOLD_LABELS - {
    'Fecha sugerida','Fecha','Grado y sección','Grado','Tema','Competencia','Capacidades',
    'Desempeño precisado','Criterio de evaluación','Criterios de evaluación','Propósito de aprendizaje','Propósito'
}

PLACEHOLDER_RE = re.compile(r'\{\{([A-Z0-9_]+)\}\}')
LABEL_RE = re.compile(r'^([^:]{1,65}):\s*(.*)$', re.S)


def _capture_run_style(paragraph: Paragraph) -> dict:
    first = paragraph.runs[0] if paragraph.runs else None
    return {
        'font_name': first.font.name if first else None,
        'font_size': first.font.size if first else None,
        'italic': first.italic if first else None,
    }


def _clear_runs(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.text = ''


def _apply_run_style(run, base: dict, *, bold: bool = False) -> None:
    run.bold = bold
    if base.get('font_name'):
        run.font.name = base['font_name']
    if base.get('font_size'):
        run.font.size = base['font_size']
    if base.get('italic') is not None:
        run.italic = base['italic']


def _label_of(text: str) -> str | None:
    m = LABEL_RE.match(text.strip())
    if not m:
        return None
    label = m.group(1).strip()
    return label if label in BOLD_LABELS else None


def _write_single_paragraph(paragraph: Paragraph, text: str, base_style: dict) -> None:
    """Write one logical block into one Word paragraph with selective bold."""
    _clear_runs(paragraph)
    text = str(text).strip()
    m = LABEL_RE.match(text)
    if m and m.group(1).strip() in BOLD_LABELS:
        label = m.group(1).strip()
        body = m.group(2)
        r1 = paragraph.add_run(label + ': ')
        _apply_run_style(r1, base_style, bold=True)
        r2 = paragraph.add_run(body)
        _apply_run_style(r2, base_style, bold=False)
    else:
        r = paragraph.add_run(text)
        _apply_run_style(r, base_style, bold=False)


def _insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    if paragraph._p.pPr is not None:
        new_p.insert(0, deepcopy(paragraph._p.pPr))
    return Paragraph(new_p, paragraph._parent)


def _logical_blocks(text: str) -> list[str]:
    """
    Convert generated text into real Word paragraphs.

    Every non-empty source line becomes a paragraph. Blank lines are represented
    by paragraph spacing instead of fake Enter characters. This is what keeps
    Inicio/Desarrollo/Cierre readable and prevents the 'all text glued together'
    defect seen in v0.2.
    """
    blocks: list[str] = []
    for line in str(text).replace('\r\n', '\n').replace('\r', '\n').split('\n'):
        line = line.strip()
        if line:
            blocks.append(line)
    return blocks or ['']


def _set_spacing(paragraph: Paragraph, label: str | None, *, is_first: bool = False) -> None:
    fmt = paragraph.paragraph_format
    if label in SECTION_LABELS:
        # Visible separation between pedagogical sub-sections, like the approved reference.
        fmt.space_before = Pt(7 if not is_first else 0)
        fmt.space_after = Pt(2)
    else:
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)


def write_multi_paragraph(paragraph: Paragraph, text: str) -> None:
    """Replace a placeholder paragraph by one or more real Word paragraphs."""
    base_style = _capture_run_style(paragraph)
    blocks = _logical_blocks(text)

    current = paragraph
    for idx, block in enumerate(blocks):
        if idx > 0:
            current = _insert_paragraph_after(current)
        _write_single_paragraph(current, block, base_style)
        _set_spacing(current, _label_of(block), is_first=(idx == 0))


def replace_para(paragraph: Paragraph, mapping: Mapping[str, str]) -> None:
    full = ''.join(r.text for r in paragraph.runs) or paragraph.text
    matches = PLACEHOLDER_RE.findall(full)
    if not matches:
        return

    stripped = full.strip()
    if len(matches) == 1 and stripped == '{{' + matches[0] + '}}':
        write_multi_paragraph(paragraph, mapping.get(matches[0], ''))
        return

    new = full
    for key in matches:
        new = new.replace('{{' + key + '}}', str(mapping.get(key, '')))
    write_multi_paragraph(paragraph, new)


def walk_cell(cell, mapping: Mapping[str, str]) -> None:
    # Snapshot because replace_para may insert sibling paragraphs.
    for paragraph in list(cell.paragraphs):
        replace_para(paragraph, mapping)
    for table in cell.tables:
        for row in table.rows:
            for child in row.cells:
                walk_cell(child, mapping)


def collect_placeholders(document: Document) -> list[str]:
    remaining: list[str] = []
    for paragraph in document.paragraphs:
        remaining.extend(PLACEHOLDER_RE.findall(paragraph.text))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                remaining.extend(PLACEHOLDER_RE.findall(cell.text))
    return sorted(set(remaining))


def fill_docx(template: str | Path, data: Mapping[str, object], output: str | Path) -> Path:
    template = Path(template)
    output = Path(output)
    document = Document(template)
    mapping = {
        k: ('' if v is None else str(v))
        for k, v in data.items()
        if not isinstance(v, (dict, list))
    }

    for paragraph in list(document.paragraphs):
        replace_para(paragraph, mapping)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                walk_cell(cell, mapping)

    for section in document.sections:
        for paragraph in list(section.header.paragraphs):
            replace_para(paragraph, mapping)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    walk_cell(cell, mapping)
        for paragraph in list(section.footer.paragraphs):
            replace_para(paragraph, mapping)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    walk_cell(cell, mapping)

    remaining = collect_placeholders(document)
    if remaining:
        raise ValueError('Unfilled placeholders: ' + ', '.join(remaining))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def fill_from_json(template: str | Path, data_file: str | Path, output: str | Path) -> Path:
    data = json.loads(Path(data_file).read_text(encoding='utf-8'))
    return fill_docx(template, data, output)
